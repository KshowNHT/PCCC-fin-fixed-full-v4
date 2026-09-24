"""Word (.docx) compliance report generator.

Direct port of src/report_generator.ts (generatePcccReport), using
python-docx as the exact functional counterpart of the `docx` npm package.
Colors, fonts, sizes and section order are preserved 1:1.
"""
import io
from datetime import date
from typing import Any, Dict, List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
PAGE_WIDTH_IN = 6.5  # usable width for an A4/Letter page with standard margins

# docx.js "size" values are in half-points; python-docx Pt() takes whole points.
def _pt(js_size: int) -> Pt:
    return Pt(js_size / 2)


def _set_cell_shading(cell, hex_color: str) -> None:
    shd = cell._tc.get_or_add_tcPr()
    shd_el = shd.makeelement(qn("w:shd"), {qn("w:fill"): hex_color})
    shd.append(shd_el)


def _run(paragraph, text: str, *, bold=False, italic=False, size=22, color=None, font=FONT):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = _pt(size)
    run.font.name = font
    run.font.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _add_paragraph(doc, text="", *, align=None, bold=False, italic=False, size=22, color=None,
                    space_after=None, bullet=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after / 20)
    if bullet:
        p.style = "List Bullet"
    if text:
        _run(p, text, bold=bold, italic=italic, size=size, color=color)
    return p


def _input_row(table, label: str, val: Any) -> None:
    row = table.add_row()
    c0, c1 = row.cells
    _run(c0.paragraphs[0], label, bold=True, size=18)
    _run(c1.paragraphs[0], str(val), size=18)


def _set_col_widths(table, weights: List[float]) -> None:
    total = sum(weights)
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(PAGE_WIDTH_IN * weights[idx] / total)


def generate_pccc_report(
    project: Dict[str, Any],
    checklist: List[Dict[str, Any]],
    conclusions: Dict[str, Any],
) -> bytes:
    is_no_appraisal = not conclusions.get("isAppraisalRequired")

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # ---- Header ----
    _add_paragraph(doc, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=24)
    _add_paragraph(doc, "Độc lập – Tự do – Hạnh phúc", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=22)
    _add_paragraph(doc, "─────────────────────────────", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=18, space_after=240)
    _add_paragraph(doc, "BẢNG TƯ VẤN AN TOÀN PHÒNG CHÁY CHỮA CHÁY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=32, color="C00000")

    name = project.get("name", "")
    floors = project.get("floors", 0)
    total_floor_area = project.get("totalFloorArea", 0)
    _add_paragraph(
        doc,
        f"{str(name).upper()} – {floors} TẦNG – {total_floor_area}m² SÀN",
        align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=26, space_after=300,
    )

    today = date.today()
    _add_paragraph(
        doc,
        f"Ngày lập: ngày {today.day} tháng {today.month} năm {today.year}",
        align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=22, space_after=400,
    )

    # ---- I. Input Data ----
    _add_paragraph(doc, "I. THÔNG TIN CÔNG TRÌNH ĐẦU VÀO", bold=True, size=24, color="1F4E79", space_after=120)

    input_table = doc.add_table(rows=1, cols=2)
    input_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = input_table.rows[0].cells
    _set_cell_shading(hdr[0], "1F4E79")
    _set_cell_shading(hdr[1], "1F4E79")
    _run(hdr[0].paragraphs[0], "THÔNG SỐ", bold=True, size=20, color="FFFFFF")
    _run(hdr[1].paragraphs[0], "GIÁ TRỊ / MÔ TẢ CHI TIẾT", bold=True, size=20, color="FFFFFF")

    _input_row(input_table, "Tên công trình", project.get("name", ""))
    _input_row(input_table, "Chủ đầu tư", project.get("investor") or "Chưa cập nhật")
    _input_row(input_table, "Địa điểm xây dựng", project.get("location") or "Chưa cập nhật")
    _input_row(
        input_table,
        "Đơn vị thiết kế / Giai đoạn",
        f"{project.get('designer') or 'Chưa cập nhật'} / {project.get('stage') or 'Thiết kế cơ sở'}",
    )
    _input_row(input_table, "Loại công trình & Công năng", project.get("type", ""))
    # FIX B21 (Bao_Cao_QA_Lan_3.md — Thấp): trước đây khi length/width = 0
    # (chưa nhập) vẫn in "0.0m (dài) × 0.0m (rộng)" như thể đã có số liệu
    # thật — trên văn bản có mục ký xác nhận tư vấn, việc này dễ gây hiểu
    # nhầm là đã đo đạc. Đổi thành "Chưa cập nhật" khi cả 2 giá trị đều 0.
    _length, _width = project.get("length") or 0, project.get("width") or 0
    _dimensions_text = f"{_length}m (dài) × {_width}m (rộng)" if (_length or _width) else "Chưa cập nhật"
    _input_row(input_table, "Kích thước hình học", _dimensions_text)
    _floor_area_val = project.get("floorArea") or 0
    _input_row(input_table, "Diện tích xây dựng tầng 1", f"{_floor_area_val} m²" if _floor_area_val else "Chưa cập nhật")
    _input_row(
        input_table,
        "Tổng diện tích sàn",
        f"~{total_floor_area} m² ({floors} tầng nổi × {project.get('floorArea', 0)} m²)" if total_floor_area else "Chưa cập nhật",
    )
    _input_row(input_table, "Quy mô tầng", f"{floors} tầng nổi, {project.get('basements', 0)} tầng hầm")
    _input_row(input_table, "Chiều cao công trình (H)", f"{project.get('height', 0)} m" if project.get("height") else "Chưa cập nhật")
    _input_row(input_table, "Chiều cao PCCC", f"{project.get('pcccHeight') or project.get('height', 0)} m" if (project.get("pcccHeight") or project.get("height")) else "Chưa cập nhật")
    _input_row(input_table, "Bậc chịu lửa dự kiến", project.get("fireRating") or "Chưa chọn")
    # FIX B21: trước đây fallback "Tầng 1 kinh doanh thương mại dịch vụ" —
    # một câu MÔ TẢ CỤ THỂ tự bịa ra khi người dùng để trống, khác hẳn cách
    # các trường khác đều fallback trung thực về "Chưa cập nhật". Đồng bộ lại.
    _input_row(
        input_table,
        "Chi tiết phần kinh doanh",
        project.get("commercialDetails") or "Chưa cập nhật",
    )
    _set_col_widths(input_table, [35, 65])

    # ---- II. Conclusion ----
    _add_paragraph(doc, "II. KẾT LUẬN VỀ THẨM DUYỆT THIẾT KẾ PCCC", bold=True, size=24, color="1F4E79", space_after=120)

    p = doc.add_paragraph()
    _run(p, "Căn cứ pháp lý chính: ", bold=True, size=22)
    _run(
        p,
        "Bộ luật PCCC, Nghị định số 105/2025/NĐ-CP ngày 15/5/2025 - Phụ lục III (Danh mục dự án, công trình "
        "thuộc diện thẩm định thiết kế về phòng cháy và chữa cháy).",
        size=22,
    )

    status_text = (
        "✓ CÔNG TRÌNH KHÔNG THUỘC DIỆN THẨM DUYỆT THIẾT KẾ PCCC BỞI CƠ QUAN CÔNG AN"
        if is_no_appraisal
        else "✗ CÔNG TRÌNH THUỘC DIỆN THẨM DUYỆT THIẾT KẾ PCCC BẮT BUỘC BỞI CƠ QUAN CÔNG AN"
    )
    _add_paragraph(
        doc, status_text, bold=True, size=24,
        color="385723" if is_no_appraisal else "C00000",
    )

    p = doc.add_paragraph()
    _run(p, "Lý do / Giải trình: ", bold=True, size=22)
    _run(p, conclusions.get("reason", ""), size=22)

    warnings = conclusions.get("warnings") or []
    if warnings:
        _add_paragraph(doc, "⚠ CẢNH BÁO / LƯU Ý BẤT THƯỜNG TRONG HỒ SƠ:", bold=True, color="C00000", size=22)
        for warn in warnings:
            _add_paragraph(doc, warn, size=22, bullet=True)

    # ---- III. Checklist Table ----
    _add_paragraph(doc, "III. CHECKLIST TỔNG HỢP KIỂM TRA RÀ SOÁT CHI TIẾT", bold=True, size=24, color="1F4E79", space_after=120)

    headers = ["STT", "HẠNG MỤC KIỂM TRA", "YÊU CẦU QUY CHUẨN", "KẾT LUẬN CHI TIẾT", "ĐÁNH GIÁ"]
    weights = [6, 22, 30, 27, 15]

    checklist_table = doc.add_table(rows=1, cols=5)
    checklist_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = checklist_table.rows[0].cells
    for i, h in enumerate(headers):
        _set_cell_shading(hdr_cells[i], "1F4E79")
        _run(hdr_cells[i].paragraphs[0], h, bold=True, color="FFFFFF", size=18)

    for item in checklist:
        result = item.get("result", "")
        status_color = "000000"
        if result == "Đạt":
            status_color = "385723"
        elif result == "Không đạt":
            status_color = "C00000"
        elif result in ("Khuyến nghị mạnh", "Cần xem xét"):
            status_color = "C05621"

        row = checklist_table.add_row()
        cells = row.cells
        _run(cells[0].paragraphs[0], str(item.get("stt", "")), size=18)

        _run(cells[1].paragraphs[0], item.get("criteria", ""), bold=True, size=18)
        ref_p = cells[1].add_paragraph()
        _run(ref_p, f"({item.get('reference', '')})", italic=True, size=16, color="595959")

        _run(cells[2].paragraphs[0], item.get("requirement", ""), size=18)
        _run(cells[3].paragraphs[0], item.get("note", ""), size=18)

        cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(cells[4].paragraphs[0], result, bold=True, color=status_color, size=18)

    _set_col_widths(checklist_table, weights)

    # ---- IV. Signatures ----
    _add_paragraph(doc, "IV. DANH MỤC CĂN CỨ VÀ KÝ XÁC NHẬN TƯ VẤN", bold=True, size=24, color="1F4E79", space_after=120)
    _add_paragraph(
        doc,
        "Các văn bản đối chiếu đã áp dụng: Nghị định số 105/2025/NĐ-CP; QCVN 06:2022/BXD Sửa đổi 1:2023; "
        "QCVN 10:2025/BCA; TCVN 13456:2022; TCVN 7435-1:2004; TCVN 7568-14:2025.",
        italic=True, size=20, space_after=240,
    )

    sig = doc.add_paragraph()
    sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(sig, "ĐƠN VỊ TƯ VẤN THẨM ĐỊNH PCCC\n\n\n\n", bold=True, size=22)
    _run(sig, "(Ký và ghi rõ họ tên, đóng dấu sản phẩm)", italic=True, size=20)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
